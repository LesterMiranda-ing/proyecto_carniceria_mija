"""Generación de reportes en Excel (.xlsx), compartida por todos los módulos.

Cada ruta de exportación solo arma encabezados + filas (datos ya formateados como
texto/número simple) y llama a `respuesta_excel`; el estilo del encabezado y el
ajuste de ancho de columnas se hacen aquí una sola vez para que todos los reportes
se vean igual.
"""

from io import BytesIO

from flask import Response
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

_RELLENO_ENCABEZADO = PatternFill(start_color="8F000D", end_color="8F000D", fill_type="solid")
_FUENTE_ENCABEZADO = Font(color="FFFFFF", bold=True)


def respuesta_excel(nombre_archivo, titulo_hoja, encabezados, filas):
    """Arma un .xlsx en memoria y lo devuelve como descarga adjunta.

    encabezados: lista de nombres de columna.
    filas: lista de tuplas/listas con los valores de cada fila, en el mismo orden.
    """
    libro = Workbook()
    hoja = libro.active
    hoja.title = titulo_hoja[:31]  # Excel limita el nombre de hoja a 31 caracteres.

    hoja.append(encabezados)
    for celda in hoja[1]:
        celda.font = _FUENTE_ENCABEZADO
        celda.fill = _RELLENO_ENCABEZADO
        celda.alignment = Alignment(horizontal="center")

    for fila in filas:
        hoja.append(list(fila))

    for columna in hoja.columns:
        ancho = max((len(str(c.value)) for c in columna if c.value is not None), default=8)
        hoja.column_dimensions[columna[0].column_letter].width = min(max(ancho + 2, 10), 40)

    hoja.freeze_panes = "A2"

    buffer = BytesIO()
    libro.save(buffer)
    buffer.seek(0)

    return Response(
        buffer.getvalue(),
        mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="{nombre_archivo}"'},
    )
